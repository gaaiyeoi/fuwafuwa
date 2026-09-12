#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""豆瓣富化器:按影片表逐条检索豆瓣条目(subject id / 海报 / 中文名 / 年份)
—— BIFF 排片工具 数据管线 M1 的一部分。

检索通道(2026-09-11 换掉)
--------------------------
旧版走网页匿名口 `movie.douban.com/j/subject_suggest`:会**静默限流**(返回 `[]`)
并最终撞 CAPTCHA,只能 60s/请求 地慢跑(250 部要 7~10 小时,见 `data/_run_2026.sh`)。
现改走官方 Frodo 口(检索 / 确认口径集中在 `tools/douban_match.py`):

* **检索** = `/api/v2/search/suggestion`(备:`search/mix_suggest_subjects`)—— 匿名稳定;
* **确认** = `/api/v2/movie/{id}` 详情 —— 用 `title / original_title / aka[]` 与本地片名对撞,
  对不上就记 `miss`,**绝不靠模糊字符串猜**(同名年份不符只给 `medium`)。

⚠ **不要**改用 `/api/v2/search/subjects`:实测跑约 100 次后返回 `403 code=103 need_login`,
而交接文档 §16.3 明确记录「豆瓣登录 / token 刷新流程未恢复」—— 拿不到 token,长跑必挂。

产物兼容
--------
`--out` 的形状与旧版**完全一致**(下游 `tools/build_films.py` 取 `subject_id` + `poster`、
`tools/fetch_posters.py` 取 `poster` URL 的档位段,都不用改);另加两个只读字段
`douban_rating` / `douban_rating_count`(刻意不与行内 `rating`=xlsx 评分撞名)。

用法:
  python enrich_douban.py --xlsx <path> --out data/enriched_douban.json [--limit N] [--delay 1.2]
  python enrich_douban.py --films-json apps/web/public/films.json --out data/enriched_douban.json

数据源(二选一):
* --xlsx        官方影片信息 xlsx —— 有 xlsx 的年份优先(列名:单元/备注/中文片名/原始片名/
                年份/评分/评价人数/国家/地区/导演)。
* --films-json  apps/web/public/films.json 目录片清单 —— 只有 PDF 产物、拿不到官方 xlsx 的年份走这条
                (2025 即属此类;字段与 xlsx 一一对应,`remark`/`rating_count` 即「备注」/「评价人数」)。
"""

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from douban_match import (  # noqa: E402  (同目录工具,需先补 sys.path)
    DoubanUnavailableError,
    dedup_queries,
    name_forms,
    poster_url_of,
    resolve_subject,
)
from frodo_client import FrodoClient  # noqa: E402


def norm(s) -> str:
    if s is None:
        return ""
    if not isinstance(s, str):
        s = str(s)
    return re.sub(r"\s+", " ", s.strip()).strip()


def match_film(client, film, delay):
    """单部影片 → `(命中字段 dict, 错误说明)`;两者必有其一为 `None`。

    检索词顺序:**中文名 → 原始片名 → 官网英文名**。中文名放最前是因为豆瓣是中文站
    (未在大陆上映的新片往往只有中文条目);英文名兜底接海外片(`aka` 里就是它)。
    """
    queries = dedup_queries(film.get("title_zh"), film.get("title_orig"), film.get("title_en"))
    forms = name_forms(film.get("title_zh"), film.get("title_orig"), film.get("title_en"))
    if not queries or not forms:
        return None, "无检索词"
    year = film["year"] if str(film.get("year") or "").isdigit() else None
    detail, confidence = resolve_subject(
        client, queries, forms, year, name_forms(film.get("director")), delay
    )
    if detail is None:
        return None, "片名对不上豆瓣条目"
    subject_id = str(detail.get("id") or "")
    rating = detail.get("rating") or {}
    value = rating.get("value")
    count = rating.get("count")
    return {
        "douban_id": subject_id,
        "douban_url": f"https://movie.douban.com/subject/{subject_id}/",
        "douban_title": detail.get("title", ""),
        "douban_orig": detail.get("original_title", ""),
        "douban_year": str(detail.get("year") or ""),
        "douban_rating": float(value) if isinstance(value, (int, float)) and value else None,
        "douban_rating_count": int(count) if isinstance(count, (int, float)) and count else None,
        "poster": poster_url_of(detail),
        "confidence": confidence,
        "score": 4.0 if confidence == "high" else 3.0,
    }, None


def load_rows(xlsx_path):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None and all(v is None for v in row):
            continue
        rows.append({
            "row": i,
            "unit": norm(row[0]), "note": norm(row[1]),
            "title_zh": norm(row[2]), "title_orig": norm(row[3]),
            "year": norm(row[4]) or "", "rating": norm(row[5]),
            "voters": norm(row[6]), "country": norm(row[7]),
            "director": norm(row[8]),
        })
    wb.close()
    return rows


def load_rows_from_json(json_path, limit=0):
    """从 apps/web/public/films.json(目录片清单)构造与 `load_rows()` **同形状**的行。

    只有 PDF 产物、拿不到官方 xlsx 的年份走这条入口(2025 即属此类)。
    字段一一对应:`remark` → 备注、`rating_count` → 评价人数;
    数值字段(year / rating / rating_count)统一转成字符串,与 xlsx 分支口径一致。
    额外带上 `title_en`(官网英文名)—— 检索时它是海外片最有效的那把钥匙。
    """
    with open(json_path, encoding="utf-8") as f:
        films = json.load(f).get("films", [])
    if limit:
        films = films[:limit]
    rows = []
    for i, film in enumerate(films, start=1):
        rows.append({
            "row": i,
            "unit": norm(film.get("unit")),
            "note": norm(film.get("remark")),
            "title_zh": norm(film.get("title_zh")),
            "title_orig": norm(film.get("title_orig")),
            "title_en": norm(film.get("title_en")),
            "year": norm(film.get("year")),
            "rating": norm(film.get("rating")),
            "voters": norm(film.get("rating_count")),
            "country": norm(film.get("country")),
            "director": norm(film.get("director")),
        })
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="官方影片信息 xlsx(与 --films-json 二选一)")
    ap.add_argument("--films-json", help="apps/web/public/films.json 目录片清单(与 --xlsx 二选一)")
    ap.add_argument("--out", default="data/enriched_douban.json")
    ap.add_argument("--limit", type=int, default=0, help="只处理前 N 部(调试用)")
    ap.add_argument("--delay", type=float, default=1.2,
                    help="每个豆瓣请求之间的间隔秒数(官方口实测 1~2s 稳定;旧版网页口要 60)")
    args = ap.parse_args()

    if bool(args.xlsx) == bool(args.films_json):
        ap.error("--xlsx 与 --films-json 必须二选一")

    if args.films_json:
        rows = load_rows_from_json(args.films_json, args.limit)
    else:
        rows = load_rows(args.xlsx)
        if args.limit:
            rows = rows[: args.limit]
    print(f"共 {len(rows)} 部影片,基础间隔 {args.delay}s")

    done = {}
    if os.path.exists(args.out):
        try:
            with open(args.out, encoding="utf-8") as f:
                for it in json.load(f):
                    if it.get("douban_id"):
                        done[it["row"]] = it
        except Exception:
            pass
    print(f"已有 {len(done)} 部完成,跳过")

    client = FrodoClient()
    stats = {"high": 0, "medium": 0, "miss": 0}
    for idx, film in enumerate(rows, 1):
        if film["row"] in done:
            continue
        try:
            hit, err = match_film(client, film, args.delay)
        except DoubanUnavailableError as exc:
            # 风控 / 断网:立即停轮。已落盘的行下次重跑自动跳过,不会把 miss 写进产物。
            print(f"\n! 豆瓣检索通道不可用,本轮停止(已完成的行已落盘,重跑可续):{exc}")
            break
        if hit:
            done[film["row"]] = {**film, **hit}
            stats[hit["confidence"]] += 1
            tag = (f"✔{hit['douban_id']} {hit['douban_title']} "
                   f"({hit['douban_year']}) conf={hit['confidence']}")
        else:
            done[film["row"]] = {**film, "douban_id": "", "confidence": "miss",
                                 "poster": "", "match_note": err}
            stats["miss"] += 1
            tag = f"✘ {err}"
        print(f"[{idx}/{len(rows)}] 行{film['row']} {film['title_zh'][:18]} -> {tag}", flush=True)
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(list(done.values()), f, ensure_ascii=False, indent=1)
        time.sleep(1.0)

    print("\n== 完成 ==", stats)
    miss = [x for x in done.values() if not x.get("douban_id")]
    if miss:
        print(f"未匹配 {len(miss)} 部,见 --out 中 confidence=miss 的行")


if __name__ == "__main__":
    main()
